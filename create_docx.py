#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador de README.docx para o Sistema LIVIA
Converte o README.md para um documento Word formatado
"""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.shared import OxmlElement, qn

def add_hyperlink(paragraph, url, text):
    """Adiciona um hyperlink ao parágrafo"""
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '0563C1')
    rPr.append(c)
    
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    
    paragraph._p.append(hyperlink)
    return hyperlink

def create_readme_docx():
    """Cria o documento README.docx"""
    doc = Document()
    
    # Configurar estilos
    styles = doc.styles
    
    # Estilo para títulos principais
    title_style = styles.add_style('Custom Title', WD_STYLE_TYPE.PARAGRAPH)
    title_font = title_style.font
    title_font.name = 'Calibri'
    title_font.size = Pt(24)
    title_font.bold = True
    title_font.color.rgb = None
    
    # Estilo para subtítulos
    subtitle_style = styles.add_style('Custom Subtitle', WD_STYLE_TYPE.PARAGRAPH)
    subtitle_font = subtitle_style.font
    subtitle_font.name = 'Calibri'
    subtitle_font.size = Pt(18)
    subtitle_font.bold = True
    
    # Estilo para seções
    section_style = styles.add_style('Custom Section', WD_STYLE_TYPE.PARAGRAPH)
    section_font = section_style.font
    section_font.name = 'Calibri'
    section_font.size = Pt(14)
    section_font.bold = True
    
    # TÍTULO PRINCIPAL
    title = doc.add_paragraph('🎓LIVIA🌼 - Sistema de Análise de Desempenho Acadêmico', style='Custom Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # SUBTÍTULO
    subtitle = doc.add_paragraph('Levantamento de Informações Vinculadas à Infrequência e Avaliação', style='Custom Subtitle')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    school = doc.add_paragraph('E.E. Anna Passamonti Balardin')
    school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    school.runs[0].italic = True
    
    # DESCRIÇÃO
    doc.add_paragraph()
    desc = doc.add_paragraph('Sistema web estático para análise de dados educacionais com funcionalidades de busca por aluno, monitoramento de notas e frequência, e geração de relatórios de risco.')
    desc.runs[0].font.size = Pt(12)
    
    # DEPLOY RÁPIDO
    doc.add_paragraph()
    doc.add_heading('🚀 Deploy Rápido', level=1)
    
    doc.add_heading('Vercel (Recomendado)', level=2)
    
    deploy_steps = [
        'Fork este repositório',
        'Conecte ao Vercel:',
        '  • Acesse vercel.com',
        '  • Importe seu repositório',
        '  • Deploy automático em segundos',
        'Acesse sua aplicação: https://seu-projeto.vercel.app'
    ]
    
    for step in deploy_steps:
        p = doc.add_paragraph(step, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # FUNCIONALIDADES
    doc.add_paragraph()
    doc.add_heading('✨ Funcionalidades', level=1)
    
    # Busca Individual
    doc.add_heading('🔍 Busca Individual', level=2)
    busca_features = [
        'Seleção por turma e série/ano',
        'Busca específica por aluno',
        'Visualização detalhada de notas e frequência',
        'Alertas automáticos para:',
        '  • 🔴 Risco de Evasão: Frequência < 75%',
        '  • ⚠️ Risco de Reprovação: Médias < 5.0'
    ]
    
    for feature in busca_features:
        p = doc.add_paragraph(feature, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # Relatórios Gerais
    doc.add_heading('📊 Relatórios Gerais', level=2)
    relatorio_features = [
        'Relatório de Notas Baixas: Estudantes com médias < 5.0 agrupados por série',
        'Relatório de Frequência Baixa: Estudantes com frequência < 75% agrupados por série',
        'Relatório Combinado: Análise integrada de ambos os riscos'
    ]
    
    for feature in relatorio_features:
        p = doc.add_paragraph(feature, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # Recursos Técnicos
    doc.add_heading('⚡ Recursos Técnicos', level=2)
    tech_features = [
        '100% Estático: Sem necessidade de servidor backend',
        'Dados Compilados: JSON pré-processado para carregamento rápido',
        'Interface Responsiva: Funciona em desktop, tablet e mobile',
        'Alertas Visuais: Destaque automático para situações de risco',
        'Tooltips Interativos: Explicações aparecem no hover do mouse',
        'Animações Críticas: Alertas piscantes para casos urgentes'
    ]
    
    for feature in tech_features:
        p = doc.add_paragraph(feature, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # ESTRUTURA DO PROJETO
    doc.add_paragraph()
    doc.add_heading('📁 Estrutura do Projeto', level=1)
    
    structure_text = """Balardin/
├── index.html              # Página principal
├── css/
│   └── style.css          # Estilos CSS com tooltips e animações
├── js/
│   └── script.js          # JavaScript principal
├── data/
│   └── alunos_compilado.json  # Dados dos alunos (compilados)
├── vercel.json            # Configuração do Vercel
├── *.pdf                  # Arquivos PDF originais
└── README.md              # Este arquivo"""
    
    structure = doc.add_paragraph(structure_text)
    structure.runs[0].font.name = 'Courier New'
    structure.runs[0].font.size = Pt(10)
    
    # COMO USAR
    doc.add_paragraph()
    doc.add_heading('🎯 Como Usar', level=1)
    
    # Busca Individual de Aluno
    doc.add_heading('1. Busca Individual de Aluno', level=2)
    busca_steps = [
        'Turma: Selecione no primeiro dropdown (ex: 1A, 2B, 3C)',
        'Série: Escolha a série/ano',
        'Aluno: Selecione o estudante específico',
        'Buscar: Clique no botão para visualizar dados completos',
        'Tooltips: Passe o mouse sobre os elementos para ver explicações'
    ]
    
    for step in busca_steps:
        p = doc.add_paragraph(step, style='List Number')
        p.runs[0].font.size = Pt(11)
    
    # Relatórios por Risco
    doc.add_heading('2. Relatórios por Risco', level=2)
    risco_items = [
        '🟡 Notas Baixas: Estudantes com risco de reprovação',
        '🔴 Frequência Baixa: Estudantes com risco de evasão',
        '🔵 Combinado: Visualização integrada dos riscos',
        '🚨 Alertas Críticos: Cards piscando vermelho para casos urgentes'
    ]
    
    for item in risco_items:
        p = doc.add_paragraph(item, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # ALERTAS VISUAIS
    doc.add_paragraph()
    doc.add_heading('🚨 Sistema de Alertas Visuais', level=1)
    
    # Situação Crítica
    doc.add_heading('🔴 Situação CRÍTICA (Animação Vermelha)', level=2)
    critico_desc = doc.add_paragraph('Quando: Aluno com notas baixas (< 5.0) E frequência baixa (< 75%)')
    critico_desc.runs[0].font.size = Pt(11)
    critico_desc.runs[0].bold = True
    
    critico_features = [
        'Fundo vermelho piscando intenso',
        'Bordas escuras destacadas',
        'Ícones animados balançando',
        'Texto "🚨 RISCO CRÍTICO ⚠️" no canto',
        'Valores críticos com animação de pulso',
        'Impossível de ignorar!'
    ]
    
    for feature in critico_features:
        p = doc.add_paragraph(feature, style='List Bullet')
        p.runs[0].font.size = Pt(10)
    
    # Risco Alto
    doc.add_heading('🟡 Risco ALTO (Animação Laranja)', level=2)
    alto_desc = doc.add_paragraph('Quando: Aluno com notas baixas OU frequência baixa (mas não ambos)')
    alto_desc.runs[0].font.size = Pt(11)
    alto_desc.runs[0].bold = True
    
    alto_features = [
        'Fundo laranja piscando moderado',
        'Ícones animados',
        'Bordas destacadas',
        'Valores críticos destacados'
    ]
    
    for feature in alto_features:
        p = doc.add_paragraph(feature, style='List Bullet')
        p.runs[0].font.size = Pt(10)
    
    # TOOLTIPS INTERATIVOS
    doc.add_paragraph()
    doc.add_heading('💬 Tooltips Interativos', level=1)
    
    tooltip_desc = doc.add_paragraph('Sistema de explicações que aparecem no hover do mouse:')
    tooltip_desc.runs[0].font.size = Pt(11)
    
    tooltip_features = [
        'Labels dos dropdowns: Explicam a função de cada seleção',
        'Botões: Mostram exatamente que resultado vão gerar',
        'Design elegante: Gradiente escuro com animação suave',
        'Posicionamento inteligente: Aparecem acima dos elementos',
        'Responsive: Adaptados para desktop e mobile'
    ]
    
    for feature in tooltip_features:
        p = doc.add_paragraph(feature, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # DADOS SUPORTADOS
    doc.add_paragraph()
    doc.add_heading('📋 Dados Suportados', level=1)
    
    doc.add_heading('Formato dos Arquivos', level=2)
    formato_items = [
        'PDFs Originais: 1A.pdf, 1B.pdf, 2A.pdf, etc.',
        'Dados Compilados: data/alunos_compilado.json'
    ]
    
    for item in formato_items:
        p = doc.add_paragraph(item, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    doc.add_heading('Critérios de Risco', level=2)
    criterios = [
        'Risco de Reprovação: Média semestral < 5.0',
        'Risco de Evasão: Frequência semestral < 75%',
        'Situação Crítica: Ambos os riscos simultaneamente'
    ]
    
    for criterio in criterios:
        p = doc.add_paragraph(criterio, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # DESENVOLVIMENTO LOCAL
    doc.add_paragraph()
    doc.add_heading('🛠️ Desenvolvimento Local', level=1)
    
    dev_code = """# Clone o repositório
git clone https://github.com/dansfisica85/Balardin.git
cd Balardin

# Sirva os arquivos estáticos
python -m http.server 8000
# ou
npx serve .

# Acesse no navegador
open http://localhost:8000"""
    
    dev_para = doc.add_paragraph(dev_code)
    dev_para.runs[0].font.name = 'Courier New'
    dev_para.runs[0].font.size = Pt(10)
    
    # TECNOLOGIAS
    doc.add_paragraph()
    doc.add_heading('📊 Tecnologias', level=1)
    
    tech_list = [
        'Frontend: HTML5, CSS3, JavaScript ES6+',
        'Dados: JSON estático pré-compilado',
        'Deploy: Vercel, GitHub Pages',
        'Design: CSS Grid, Flexbox, Responsive Design',
        'Animações: CSS Keyframes, Transform, Transition',
        'Tooltips: CSS Hover States, JavaScript',
        'Performance: Optimized assets, cache headers'
    ]
    
    for tech in tech_list:
        p = doc.add_paragraph(tech, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # SOLUÇÃO DE PROBLEMAS
    doc.add_paragraph()
    doc.add_heading('🐛 Solução de Problemas', level=1)
    
    doc.add_heading('Dados não carregam', level=2)
    dados_solucoes = [
        'Verifique se data/alunos_compilado.json existe',
        'Abra o console do navegador (F12) para ver erros',
        'Verifique se está servindo via HTTP (não file://)'
    ]
    
    for solucao in dados_solucoes:
        p = doc.add_paragraph(solucao, style='List Number')
        p.runs[0].font.size = Pt(10)
    
    doc.add_heading('Tooltips não aparecem', level=2)
    tooltip_solucoes = [
        'Verifique se css/style.css está carregando',
        'Teste em navegador diferente',
        'Limpe o cache do navegador (Ctrl+F5)'
    ]
    
    for solucao in tooltip_solucoes:
        p = doc.add_paragraph(solucao, style='List Number')
        p.runs[0].font.size = Pt(10)
    
    doc.add_heading('Deploy no Vercel', level=2)
    deploy_solucoes = [
        'Confirme que o repositório está público ou conectado',
        'Verifique se vercel.json está na raiz',
        'Confira os logs de build no dashboard do Vercel'
    ]
    
    for solucao in deploy_solucoes:
        p = doc.add_paragraph(solucao, style='List Number')
        p.runs[0].font.size = Pt(10)
    
    # LICENÇA E SUPORTE
    doc.add_paragraph()
    doc.add_heading('📄 Licença', level=1)
    licenca = doc.add_paragraph('Sistema desenvolvido para análise educacional. Livre para uso e modificação conforme necessidades específicas da instituição.')
    licenca.runs[0].font.size = Pt(11)
    
    doc.add_paragraph()
    doc.add_heading('🆘 Suporte', level=1)
    suporte_items = [
        '📖 Consulte este README',
        '🔍 Verifique o console do navegador (F12)',
        '🌐 Teste o deploy no Vercel',
        '📧 Abra uma issue no GitHub'
    ]
    
    for item in suporte_items:
        p = doc.add_paragraph(item, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # FOOTER
    doc.add_paragraph()
    footer = doc.add_paragraph('Desenvolvido com ❤️ para educação de qualidade')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(12)
    
    # Adicionar quebra de página
    doc.add_page_break()
    
    # PÁGINA DE MELHORIAS RECENTES
    doc.add_heading('🚀 Últimas Melhorias Implementadas', level=1)
    
    melhorias_data = doc.add_paragraph('Agosto 2025 - Versão com Tooltips e Alertas Críticos')
    melhorias_data.runs[0].bold = True
    melhorias_data.runs[0].font.size = Pt(12)
    
    melhorias_list = [
        '🎭 Tooltips hover elegantes substituindo descrições fixas',
        '🚨 Sistema de alertas críticos com animações piscantes',
        '🌊 Sombras otimizadas para scroll suave',
        '🖼️ Molduras escuras separando seções visualmente',
        '📱 Interface responsiva aprimorada',
        '⚡ Performance otimizada para carregamento rápido',
        '🎨 Design moderno e profissional'
    ]
    
    for melhoria in melhorias_list:
        p = doc.add_paragraph(melhoria, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # Salvar documento
    doc.save('/workspaces/Balardin/README_LIVIA.docx')
    print("✅ README_LIVIA.docx criado com sucesso!")
    print("📁 Localização: /workspaces/Balardin/README_LIVIA.docx")

if __name__ == "__main__":
    create_readme_docx()
